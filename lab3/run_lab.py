from __future__ import annotations

import csv
import math
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Tuple

from PIL import Image as PILImage, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from openpyxl import load_workbook

OUT = Path("outputs")
SEQUENCES = {
    "S1": "ATTGCCATT", "S2": "ATGGCCATT", "S3": "ATCCAATTTT",
    "S4": "ATCTTCTT", "S5": "ACTGACC",
}

@dataclass
class Node:
    name: str
    left: "Node | None" = None
    right: "Node | None" = None
    left_length: float = 0.0
    right_length: float = 0.0

def global_alignment(a: str, b: str) -> Tuple[str, str]:

    n, m = len(a), len(b)
    score = [[0] * (m + 1) for _ in range(n + 1)]
    for i in range(1, n + 1): score[i][0] = -i
    for j in range(1, m + 1): score[0][j] = -j
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            score[i][j] = max(score[i-1][j-1] + (1 if a[i-1] == b[j-1] else -1),
                              score[i-1][j] - 1, score[i][j-1] - 1)
    x, y, i, j = [], [], n, m
    while i or j:
        if i and j and score[i][j] == score[i-1][j-1] + (1 if a[i-1] == b[j-1] else -1):
            x.append(a[i-1]); y.append(b[j-1]); i -= 1; j -= 1
        elif i and score[i][j] == score[i-1][j] - 1:
            x.append(a[i-1]); y.append("-"); i -= 1
        else:
            x.append("-"); y.append(b[j-1]); j -= 1
    return "".join(reversed(x)), "".join(reversed(y))

def jukes_cantor_distance(a: str, b: str) -> float:
    x, y = global_alignment(a, b)
    p = sum(c1 != c2 for c1, c2 in zip(x, y)) / len(x)
    if p >= 0.75:
        raise ValueError(f"p={p:.3f}; JC69 no esta definida cuando p >= 0.75")
    return -0.75 * math.log(1 - (4 / 3) * p)

def sequence_matrix(seqs: Dict[str, str]) -> Dict[Tuple[str, str], float]:
    names = list(seqs)
    d = {(x, x): 0.0 for x in names}
    for i, x in enumerate(names):
        for y in names[i+1:]: d[x, y] = d[y, x] = jukes_cantor_distance(seqs[x], seqs[y])
    return d

def surname_distance(a: str, b: str) -> float:
    
    def normal(text: str) -> str:
        return "".join(c for c in unicodedata.normalize("NFD", text.upper()) if unicodedata.category(c) != "Mn")
    a, b = normal(a), normal(b)
    previous = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        current = [i]
        for j, cb in enumerate(b, 1): current.append(min(current[-1]+1, previous[j]+1, previous[j-1] + (ca != cb)))
        previous = current
    return previous[-1] / max(len(a), len(b), 1)

def analyze_surnames() -> None:
    path = OUT / "apellidos_clase.xlsx"
    if not path.exists():
        print("No se encontro outputs/apellidos_clase.xlsx; se omitio el analisis de apellidos.")
        return
    sheet = load_workbook(path, data_only=True)["Apellidos"]
    rows = [row for row in sheet.iter_rows(min_row=6, values_only=True) if row[0] is not None]
    if len(rows) < 2: return
    labels = [str(row[1]).strip() for row in rows]
    surnames = [f"{str(row[2]).strip()} {str(row[3]).strip()}" for row in rows]
    d = {(x, x): 0.0 for x in labels}
    for i, x in enumerate(labels):
        for j in range(i+1, len(labels)):
            d[x, labels[j]] = d[labels[j], x] = surname_distance(surnames[i], surnames[j])
    up, nj = upgma(labels, d), neighbor_joining(labels, d)
    draw_tree(up, OUT / "apellidos_upgma.png", "Apellidos: UPGMA")
    draw_tree(nj, OUT / "apellidos_neighbor_joining.png", "Apellidos: Neighbor Joining")
    (OUT / "apellidos_upgma.nwk").write_text(newick(up) + ";\n", encoding="utf-8")
    (OUT / "apellidos_neighbor_joining.nwk").write_text(newick(nj) + ";\n", encoding="utf-8")

def upgma(names: List[str], source: Dict[Tuple[str, str], float]) -> Node:
    clusters = {x: ([x], Node(x), 0.0) for x in names}
    while len(clusters) > 1:
        a, b = min(((a, b) for a in clusters for b in clusters if a < b),
                   key=lambda z: sum(source[x, y] for x in clusters[z[0]][0] for y in clusters[z[1]][0]) /
                                 (len(clusters[z[0]][0]) * len(clusters[z[1]][0])))
        leaves_a, node_a, height_a = clusters.pop(a); leaves_b, node_b, height_b = clusters.pop(b)
        height = sum(source[x, y] for x in leaves_a for y in leaves_b) / (len(leaves_a) * len(leaves_b)) / 2
        parent = Node(f"({a},{b})", node_a, node_b, height - height_a, height - height_b)
        clusters[parent.name] = (leaves_a + leaves_b, parent, height)
    return next(iter(clusters.values()))[1]

def neighbor_joining(names: List[str], d: Dict[Tuple[str, str], float]) -> Node:
    nodes = {x: Node(x) for x in names}; active = list(names); counter = 1
    dist = dict(d)
    while len(active) > 2:
        n = len(active); r = {x: sum(dist[x, y] for y in active if y != x) for x in active}
        a, b = min(((x, y) for i, x in enumerate(active) for y in active[i+1:]),
                   key=lambda z: (n - 2) * dist[z] - r[z[0]] - r[z[1]])
        limb_a = max(0.0, 0.5 * (dist[a, b] + (r[a] - r[b]) / (n - 2)))
        limb_b = max(0.0, dist[a, b] - limb_a)
        u = f"N{counter}"; counter += 1
        nodes[u] = Node(u, nodes[a], nodes[b], limb_a, limb_b)
        for x in active:
            if x not in (a, b): dist[u, x] = dist[x, u] = 0.5 * (dist[a, x] + dist[b, x] - dist[a, b])
        active = [x for x in active if x not in (a, b)] + [u]
    a, b = active
    return Node("root", nodes[a], nodes[b], dist[a, b] / 2, dist[a, b] / 2)

def newick(node: Node) -> str:
    if not node.left: return node.name
    return f"({newick(node.left)}:{node.left_length:.4f},{newick(node.right)}:{node.right_length:.4f}){node.name}"

def draw_tree(root: Node, filename: Path, title: str) -> None:
    leaves: List[Node] = []
    def get_leaves(n: Node):
        if n.left is None: leaves.append(n)
        else: get_leaves(n.left); get_leaves(n.right)
    get_leaves(root); y = {id(n): i for i, n in enumerate(leaves)}
    max_x = 0.0
    def depth(n: Node, x: float = 0.0) -> None:
        nonlocal max_x
        max_x = max(max_x, x)
        if n.left: depth(n.left, x + n.left_length); depth(n.right, x + n.right_length)
    depth(root)
    image = PILImage.new("RGB", (1100, max(260, 110 + 85 * len(leaves))), "white")
    canvas = ImageDraw.Draw(image); font = ImageFont.load_default()
    def xy(x: float, yy: float) -> Tuple[int, int]:
        return 100 + int(800 * x / max(max_x, .01)), 65 + int(75 * yy)
    def place(n: Node, x: float) -> float:
        if n.left is None:
            px, py = xy(x, y[id(n)]); canvas.text((px + 8, py - 6), n.name, fill="black", font=font); return y[id(n)]
        yl = place(n.left, x + n.left_length); yr = place(n.right, x + n.right_length)
        px, _ = xy(x, 0); lx, ly = xy(x + n.left_length, yl); rx, ry = xy(x + n.right_length, yr)
        canvas.line((px, ly, px, ry), fill="black", width=2); canvas.line((px, ly, lx, ly), fill="black", width=2); canvas.line((px, ry, rx, ry), fill="black", width=2)
        return (yl + yr) / 2
    place(root, 0)
    canvas.text((25, 20), title, fill="black", font=font)
    canvas.text((100, image.height - 25), "Distancia evolutiva (longitud horizontal proporcional)", fill="black", font=font)
    image.save(filename)

def save_matrix(names: List[str], d: Dict[Tuple[str, str], float]) -> None:
    with (OUT / "matriz_distancias.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow([""] + names)
        for x in names: w.writerow([x] + [f"{d[x, y]:.6f}" for y in names])

def main() -> None:
    OUT.mkdir(exist_ok=True)
    names = list(SEQUENCES); d = sequence_matrix(SEQUENCES)
    save_matrix(names, d)
    up = upgma(names, d); nj = neighbor_joining(names, d)
    (OUT / "arbol_upgma.nwk").write_text(newick(up) + ";\n", encoding="utf-8")
    (OUT / "arbol_neighbor_joining.nwk").write_text(newick(nj) + ";\n", encoding="utf-8")
    draw_tree(up, OUT / "arbol_upgma.png", "UPGMA")
    draw_tree(nj, OUT / "arbol_neighbor_joining.png", "Neighbor Joining")
    analyze_surnames()
    build_report(names, d)
    print("Listo. Resultados en outputs/")

if __name__ == "__main__": main()
